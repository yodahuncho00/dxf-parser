import csv
from datetime import datetime
import logging
from pathlib import Path
import re
from typing import Any, Iterator, List

import ezdxf
from ezdxf.lldxf.const import DXFStructureError
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

INPUT_LAYER = "GT_1"
OUTPUT_FILE = "pipe_bom_export.csv"
MTO_DIR = "MTO"
MTO_EXCEL_FILE = "MTO_Pipes_AG.xlsx"
PIPE_PREFIX = "PIPE B36.10"
CSV_HEADERS = [
    "Drawing Number",
    "Rev.",
    "Date",
    "Reason for issue",
    "Area",
    "Pipe description",
    "material",
    "diam",
    "Sch.",
    "material code",
    "Qty.",
]
EXCEL_HEADERS = [
    "Drawing Number",
    "Rev.",
    "Date",
    "Reason for issue",
    "Area",
    "Pipe description",
    "material",
    "diameter",
    "Sch.",
    "material code",
    "Qty.",
    "ISO Base",
    "Revision Status",
]

SCHEDULE_PATTERN = re.compile(r"\s*-?\s*Sch\.\s*([^\s,;]+)", re.IGNORECASE)
SCHEDULE_ONLY_PATTERN = re.compile(r"^\s*Sch\.\s*(.+?)\s*$", re.IGNORECASE)
QUANTITY_METERS_PATTERN = re.compile(r"\s*M\s*$", re.IGNORECASE)
MATERIAL_PATTERN = re.compile(r"\b(A\d{2,4})\b", re.IGNORECASE)
DATE_REGEX = re.compile(r"^\d{2}[./-]\d{2}[./-]\d{4}$|^\d{4}[./-]\d{2}[./-]\d{2}$")
DATE_FORMATS = [
    "%d.%m.%Y",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%Y-%m-%d",
    "%Y.%m.%d",
]
STATUS_KEYWORDS = {
    "ISSUED",
    "IFC",
    "IFD",
    "IFA",
    "IFR",
    "CONSTRUCTION",
    "DESIGN",
    "REVIEW",
    "APPROVAL",
    "BID",
    "TENDER",
    "HAZOP",
    "COMMENT",
    "FEED",
    "AS-BUILT",
    "RECORD",
    "PRELIMINARY",
}


def parse_date(text: str) -> datetime | None:
    """Attempt to parse date string into a datetime object."""
    if not DATE_REGEX.match(text):
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def extract_latest_revision_record(all_text_entities: List[tuple[str, str]]) -> tuple[str, str]:
    """Scan all modelspace TEXT entities for revision records and return (date, reason_for_issue)
    corresponding to the latest date.
    """
    candidates: list[tuple[datetime, str, str]] = []

    for i, (text, _) in enumerate(all_text_entities):
        t = text.strip()
        dt = parse_date(t)
        if dt is not None:
            # Look for status in nearby entities (forward first, then backward)
            status_val = ""
            for j in range(i + 1, min(i + 7, len(all_text_entities))):
                nt = all_text_entities[j][0].strip()
                if any(kw in nt.upper() for kw in STATUS_KEYWORDS):
                    status_val = nt
                    break
            if not status_val:
                for j in range(i - 1, max(-1, i - 6), -1):
                    nt = all_text_entities[j][0].strip()
                    if any(kw in nt.upper() for kw in STATUS_KEYWORDS):
                        status_val = nt
                        break

            candidates.append((dt, t, status_val))

    if not candidates:
        return "", ""

    # Sort by date ascending, prioritizing candidates with a non-empty status in case of same date
    candidates.sort(key=lambda c: (c[0], bool(c[2])))
    latest = candidates[-1]
    return latest[1], latest[2]


def clean_dwg_name(stem: str) -> str:
    """Strip duplicate download/copy suffixes such as ' (1)', ' (2)', ' (copy)' from drawing name."""
    cleaned = re.sub(r"\s*\(\d+\)$", "", stem).strip()
    cleaned = re.sub(r"\s*\(copy(?:\s+\d+)?\)$", "", cleaned, flags=re.IGNORECASE).strip()
    return cleaned


def extract_rev(dwg_name: str) -> str:
    """Extract revision number from drawing name, defaulting to '0'."""
    match = re.search(r"_Rev(\w+)", dwg_name, re.IGNORECASE)
    if match:
        return match.group(1)
    match = re.search(r"_R(\d+)", dwg_name, re.IGNORECASE)
    if match:
        return match.group(1)
    return "0"


def extract_area(dwg_name: str) -> str:
    """Extract 4-digit Area code from drawing name (e.g., '8906' from '...-8906-8960-...', '8631' from '...-860U-8631-...')."""
    match = re.search(r"-(\d{4})-\d{4}-", dwg_name)
    if match:
        return match.group(1)
    match = re.search(r"-(\d{4})-", dwg_name)
    if match:
        return match.group(1)
    return ""


def split_schedule(description: str) -> tuple[str, str]:
    """Return a description without its schedule and the schedule without 'Sch.'."""
    match = SCHEDULE_PATTERN.search(description)
    if not match:
        return description, ""

    schedule = match.group(1)
    description_without_schedule = (
        description[: match.start()] + description[match.end() :]
    ).strip(" -")
    return description_without_schedule, schedule


def schedule_value(text: str) -> str | None:
    """Return the value of a standalone schedule TEXT, without the 'Sch.' prefix."""
    match = SCHEDULE_ONLY_PATTERN.match(text)
    return match.group(1) if match else None


def normalize_quantity(quantity: str) -> str:
    """Remove a trailing metres unit from the exported quantity."""
    return QUANTITY_METERS_PATTERN.sub("", quantity).strip()


def extract_material(description: str) -> str:
    """Extract the first ASTM-style material code, such as A333, from a description."""
    match = MATERIAL_PATTERN.search(description)
    return match.group(1).upper() if match else ""


def extract_modelspace_text_entities(filepath: Path) -> List[tuple[str, str]]:
    """Fast, low-level line-by-line parsing of ASCII DXF file to extract TEXT entities in Model Space.
    Returns a list of (text_value, layer_name) tuples.
    """
    text_entities = []
    in_entities_section = False

    current_entity_type = None
    current_text = ""
    current_layer = ""
    is_paperspace = False
    in_model_layout = True

    try:
        with open(filepath, "r", encoding="utf-8", errors="surrogateescape") as f:
            lines = f.readlines()
    except Exception:
        with open(filepath, "r", encoding="cp1252", errors="surrogateescape") as f:
            lines = f.readlines()

    it = iter(lines)
    while True:
        try:
            line_code = next(it)
        except StopIteration:
            break
        try:
            line_val = next(it)
        except StopIteration:
            break

        code_str = line_code.strip()
        if not code_str:
            continue
        try:
            code = int(code_str)
        except ValueError:
            continue

        val = line_val.strip()

        if code == 0:
            # Save the completed TEXT entity if we were building one
            if in_entities_section and current_entity_type == "TEXT":
                if not is_paperspace and in_model_layout:
                    text_entities.append((current_text.strip(), current_layer))

            # Reset entity fields for the next entity
            current_entity_type = None
            current_text = ""
            current_layer = ""
            is_paperspace = False
            in_model_layout = True

            # Section boundaries or new entities
            if val == "SECTION":
                try:
                    next_code_line = next(it)
                    next_val_line = next(it)
                    next_code = int(next_code_line.strip())
                    next_val = next_val_line.strip()
                    if next_code == 2 and next_val == "ENTITIES":
                        in_entities_section = True
                except (StopIteration, ValueError):
                    pass
            elif val == "ENDSEC":
                in_entities_section = False
            elif in_entities_section:
                current_entity_type = val
        else:
            if in_entities_section and current_entity_type == "TEXT":
                if code == 1:
                    current_text = val
                elif code == 8:
                    current_layer = val
                elif code == 67:
                    is_paperspace = val == "1"
                elif code == 410:
                    in_model_layout = val.upper() == "MODEL"

    return text_entities


def extract_pipe_rows(dxf_path: Path) -> Iterator[List[str]]:
    """Extract pipe rows along with revision Date and Status from the DXF file."""
    all_text_entities = []
    use_fallback = False
    try:
        all_text_entities = extract_modelspace_text_entities(dxf_path)
        if not all_text_entities:
            use_fallback = True
    except Exception as exc:
        logging.debug(
            "Fast parser failed on %s: %s. Falling back to ezdxf.", dxf_path.name, exc
        )
        use_fallback = True

    if use_fallback:
        try:
            document = ezdxf.readfile(dxf_path)
            all_text_entities = [
                (entity.dxf.text, entity.dxf.layer)
                for entity in document.modelspace()
                if entity.dxftype() == "TEXT"
            ]
        except (IOError, OSError, DXFStructureError, ezdxf.DXFError) as exc:
            logging.warning("Skipping invalid DXF %s: %s", dxf_path.name, exc)
            return

    # Extract date and reason for issue from the latest revision record
    date_val, status_val = extract_latest_revision_record(all_text_entities)

    # Extract GT_1 TEXT values for the pipe BOM parser
    values = [
        text.strip() for text, layer in all_text_entities if layer == INPUT_LAYER
    ]

    dwg_name = clean_dwg_name(dxf_path.stem)
    rev_val = extract_rev(dwg_name)
    area_val = extract_area(dwg_name)

    index = 0
    while index < len(values):
        if values[index].upper().startswith(PIPE_PREFIX):
            if index + 3 >= len(values):
                logging.warning(
                    "Incomplete pipe record in %s after TEXT value %r",
                    dxf_path.name,
                    values[index],
                )
                index += 1
                continue

            description = values[index]
            standalone_schedule = schedule_value(values[index + 1])

            if standalone_schedule is not None:
                if index + 4 >= len(values):
                    logging.warning(
                        "Incomplete pipe record in %s after TEXT value %r",
                        dxf_path.name,
                        description,
                    )
                    index += 1
                    continue
                schedule = standalone_schedule
                diameter, material_code, quantity = values[index + 2 : index + 5]
                index += 5
            else:
                description, schedule = split_schedule(description)
                diameter, material_code, quantity = values[index + 1 : index + 4]
                index += 4

            yield [
                dwg_name,
                rev_val,
                date_val,
                status_val,
                area_val,
                description,
                extract_material(description),
                diameter,
                schedule,
                material_code,
                normalize_quantity(quantity),
            ]
        else:
            index += 1


def process_dxf_file(dxf_path: Path) -> List[List[str]]:
    """Worker function for multiprocessing pool to parse a single DXF file."""
    return list(extract_pipe_rows(dxf_path))


def append_to_mto_excel(all_rows: List[List[str]], excel_path: Path) -> None:
    """Append extracted pipe BOM rows to MTO.xlsx with ISO Base, dynamic Revision Status,
    expanded table range, and filter-aware summary formulas.
    """
    wb: openpyxl.Workbook
    ws: openpyxl.worksheet.worksheet.Worksheet

    if excel_path.exists():
        try:
            wb = openpyxl.load_workbook(excel_path)
        except PermissionError:
            logging.error(
                "Could not open '%s' because it is locked by another program (e.g. Excel). "
                "Please close '%s' and run the script again.",
                excel_path.name,
                excel_path.name,
            )
            return
        ws = wb["Total"] if "Total" in wb.sheetnames else wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Total"
        # Write headers at row 2 (columns B to N)
        for col_idx, header in enumerate(EXCEL_HEADERS, start=2):
            ws.cell(row=2, column=col_idx, value=header)

    # Read existing data rows (cols 2 to 12) and deduplicate
    existing_rows: List[List[Any]] = []
    incoming_dwgs = {clean_dwg_name(r[0]) for r in all_rows} if all_rows else set()

    for r in range(3, ws.max_row + 1):
        val_b = ws.cell(row=r, column=2).value
        if val_b is not None and not str(val_b).startswith("="):
            cleaned_dwg = clean_dwg_name(str(val_b))
            # If this drawing is being newly supplied by all_rows, replace it (don't keep old duplicate)
            if cleaned_dwg in incoming_dwgs:
                continue
            row_vals = [ws.cell(row=r, column=c).value for c in range(2, 13)]
            row_vals[0] = cleaned_dwg
            existing_rows.append(row_vals)

    # Deduplicate existing rows to remove any previous duplicate runs
    seen = set()
    deduped_existing: List[List[Any]] = []
    for erow in existing_rows:
        key = tuple(str(x) for x in erow)
        if key not in seen:
            seen.add(key)
            deduped_existing.append(erow)

    # Clear worksheet data area from row 3 downwards
    max_clean_row = max(ws.max_row + 5, len(deduped_existing) + len(all_rows) + 10)
    for r in range(3, max_clean_row + 1):
        for c in range(1, 16):
            ws.cell(row=r, column=c).value = None
            ws.cell(row=r, column=c).border = Border()
            ws.cell(row=r, column=c).fill = PatternFill(fill_type=None)
            ws.cell(row=r, column=c).font = Font(name="Calibri", size=11)

    # 1. Write back existing retained rows
    next_row = 3
    for erow in deduped_existing:
        dwg_name, rev_val, date_val, status_val, area_val, desc, mat, diam, sch, mat_code, qty = erow
        
        # Clean Area
        extracted_area = extract_area(str(dwg_name))
        typed_area = int(extracted_area) if extracted_area and extracted_area.isdigit() else (int(area_val) if str(area_val).isdigit() else area_val)
        
        # Clean Rev
        typed_rev = int(str(rev_val)) if str(rev_val).isdigit() else rev_val
        
        # Clean Qty
        typed_qty = qty
        if qty is not None and not str(qty).startswith("="):
            qstr = str(qty).replace(",", ".").strip()
            try:
                typed_qty = float(qstr) if "." in qstr else int(qstr)
            except ValueError:
                typed_qty = qty

        iso_base_formula = f'=IFERROR(LEFT(B{next_row},SEARCH("_Rev",B{next_row})-1),B{next_row})'

        row_values = [
            (2, dwg_name, Alignment(horizontal="left", vertical="center")),
            (3, typed_rev, Alignment(horizontal="center", vertical="center")),
            (4, date_val, Alignment(horizontal="center", vertical="center")),
            (5, status_val, Alignment(horizontal="left", vertical="center")),
            (6, typed_area, Alignment(horizontal="center", vertical="center")),
            (7, desc, Alignment(horizontal="left", vertical="center")),
            (8, mat, Alignment(horizontal="left", vertical="center")),
            (9, diam, Alignment(horizontal="right", vertical="center")),
            (10, sch, Alignment(horizontal="center", vertical="center")),
            (11, mat_code, Alignment(horizontal="left", vertical="center")),
            (12, typed_qty, Alignment(horizontal="right", vertical="center")),
            (13, iso_base_formula, Alignment(horizontal="left", vertical="center")),
            (14, "", Alignment(horizontal="left", vertical="center")),
        ]
        for col_idx, val, align in row_values:
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            cell.alignment = align
            cell.font = Font(name="Calibri", size=11)
        next_row += 1

    # 2. Append newly extracted rows with green highlight
    green_appended_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for row_data in all_rows:
        (
            dwg_name,
            rev_val,
            date_val,
            status_val,
            area_val,
            description,
            material,
            diameter,
            schedule,
            material_code,
            quantity,
        ) = row_data

        clean_name = clean_dwg_name(dwg_name)

        # Numeric conversions
        typed_qty: Any
        clean_qty = str(quantity).replace(",", ".").strip()
        try:
            typed_qty = float(clean_qty) if "." in clean_qty else int(clean_qty)
        except ValueError:
            typed_qty = quantity

        typed_rev: Any
        try:
            typed_rev = int(rev_val)
        except ValueError:
            typed_rev = rev_val

        extracted_area = extract_area(clean_name)
        typed_area = int(extracted_area) if extracted_area and extracted_area.isdigit() else (int(area_val) if str(area_val).isdigit() else area_val)

        typed_diam: Any
        clean_diam = str(diameter).replace(",", ".").strip()
        try:
            typed_diam = float(clean_diam) if "." in clean_diam else int(clean_diam)
        except ValueError:
            typed_diam = diameter

        iso_base_formula = f'=IFERROR(LEFT(B{next_row},SEARCH("_Rev",B{next_row})-1),B{next_row})'

        row_values = [
            (2, clean_name, Alignment(horizontal="left", vertical="center")),
            (3, typed_rev, Alignment(horizontal="center", vertical="center")),
            (4, date_val, Alignment(horizontal="center", vertical="center")),
            (5, status_val, Alignment(horizontal="left", vertical="center")),
            (6, typed_area, Alignment(horizontal="center", vertical="center")),
            (7, description, Alignment(horizontal="left", vertical="center")),
            (8, material, Alignment(horizontal="left", vertical="center")),
            (9, typed_diam, Alignment(horizontal="right", vertical="center")),
            (10, schedule, Alignment(horizontal="center", vertical="center")),
            (11, material_code, Alignment(horizontal="left", vertical="center")),
            (12, typed_qty, Alignment(horizontal="right", vertical="center")),
            (13, iso_base_formula, Alignment(horizontal="left", vertical="center")),
            (14, "", Alignment(horizontal="left", vertical="center")),
        ]

        for col_idx, val, align in row_values:
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            cell.alignment = align
            cell.font = Font(name="Calibri", size=11)
            cell.fill = green_appended_fill
            cell.border = thin_border

        next_row += 1

    new_last_row = next_row - 1

    if new_last_row >= 3:
        # Normalize data types and populate standard formulas across all rows
        for r in range(3, new_last_row + 1):
            # Clean Drawing Number and ensure Area is accurate 4-digit code
            dwg_str = str(ws.cell(row=r, column=2).value or "")
            cell_area = ws.cell(row=r, column=6)
            extracted = extract_area(dwg_str)
            if extracted:
                try:
                    cell_area.value = int(extracted)
                except ValueError:
                    cell_area.value = extracted
            elif cell_area.value is not None:
                area_str = str(cell_area.value).strip()
                try:
                    cell_area.value = int(area_str)
                except ValueError:
                    cell_area.value = area_str

            # Normalize Rev
            cell_rev = ws.cell(row=r, column=3)
            if cell_rev.value is not None:
                rev_str = str(cell_rev.value).strip()
                try:
                    cell_rev.value = int(rev_str)
                except ValueError:
                    cell_rev.value = rev_str

            # Normalize Qty
            cell_qty = ws.cell(row=r, column=12)
            if cell_qty.value is not None and not str(cell_qty.value).startswith("="):
                qty_str = str(cell_qty.value).replace(",", ".").strip()
                try:
                    cell_qty.value = float(qty_str) if "." in qty_str else int(qty_str)
                except ValueError:
                    pass

            cell_m = ws.cell(row=r, column=13)
            cell_n = ws.cell(row=r, column=14)
            cell_m.value = f'=IFERROR(LEFT(B{r},SEARCH("_Rev",B{r})-1),B{r})'
            cell_n.value = (
                f'=IF(B{r}="","",IF(ISNUMBER(C{r}),'
                f'IF(COUNTIFS($M$3:$M${new_last_row},M{r},$C$3:$C${new_last_row},">"&C{r})=0,"CURRENT","SUPERSEDED"),"CURRENT"))'
            )

        # Update or create clean table (remove any queryTable artifacts)
        if "pipe_bom_export" in ws.tables:
            del ws.tables["pipe_bom_export"]

        tbl = Table(displayName="pipe_bom_export", ref=f"B2:N{new_last_row}")
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium7",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tbl)

        # Do not set worksheet-level auto_filter to avoid conflicting with table autoFilter
        ws.auto_filter.ref = None

        # Summary Row immediately after data rows
        summary_row = new_last_row + 1

        # i) Unique count of displayed Drawing Numbers (using _xlfn prefix for OpenXML Future Functions)
        dwg_unique_formula = (
            f'=_xlfn.IFERROR(_xlfn.ROWS(_xlfn.UNIQUE(_xlfn._xlws.FILTER(B3:B{new_last_row}, '
            f'(B3:B{new_last_row}<>"") * (SUBTOTAL(103, OFFSET(B3, ROW(B3:B{new_last_row})-ROW(B3), 0, 1, 1))=1)))), 0)'
        )
        # ii) Unique count of displayed Areas
        area_unique_formula = (
            f'=_xlfn.IFERROR(_xlfn.ROWS(_xlfn.UNIQUE(_xlfn._xlws.FILTER(F3:F{new_last_row}, '
            f'(F3:F{new_last_row}<>"") * (SUBTOTAL(103, OFFSET(F3, ROW(F3:F{new_last_row})-ROW(F3), 0, 1, 1))=1)))), 0)'
        )
        # iii) Sum of displayed Qty.
        qty_sum_formula = f"=SUBTOTAL(109, L3:L{new_last_row})"

        ws.cell(row=summary_row, column=2, value=dwg_unique_formula).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.cell(row=summary_row, column=6, value=area_unique_formula).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        ws.cell(row=summary_row, column=12, value=qty_sum_formula).alignment = Alignment(
            horizontal="right", vertical="center"
        )

        bold_font = Font(name="Calibri", size=11, bold=True)
        summary_border = Border(
            top=Side(style="thin", color="000000"),
            bottom=Side(style="double", color="000000"),
        )
        for c in range(2, 15):
            cell = ws.cell(row=summary_row, column=c)
            cell.font = bold_font
            cell.border = summary_border

        # Prune leftover empty cell objects beyond summary_row so dimensions match exactly
        leftover_coords = [coord for coord, cell in ws._cells.items() if cell.row > summary_row]
        for coord in leftover_coords:
            del ws._cells[coord]

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = col[0].column_letter
        col_idx = col[0].column
        if 2 <= col_idx <= 14:
            max_len = 0
            for cell in col:
                if cell.row > new_last_row + 1:
                    continue
                val_str = str(cell.value if cell.value is not None else "")
                if val_str.startswith("="):
                    continue
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    try:
        excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(excel_path)
        logging.info("Saved updated MTO workbook to %s", excel_path)
    except PermissionError:
        logging.error(
            "Could not save '%s' because it is locked by another program (e.g. Excel). "
            "Please close '%s' and run the script again.",
            excel_path.name,
            excel_path.name,
        )


def main() -> None:
    folder = Path.cwd()
    output_path = folder / OUTPUT_FILE
    mto_excel_path = folder / MTO_DIR / MTO_EXCEL_FILE
    dxf_files = sorted(
        path
        for path in folder.rglob("*")
        if path.is_file() and path.suffix.lower() == ".dxf"
    )
    row_count = 0
    empty_dxf_files: list[Path] = []
    all_rows: list[list[str]] = []

    import multiprocessing

    if dxf_files:
        with output_path.open("w", newline="", encoding="utf-8-sig") as csv_file:
            writer = csv.writer(csv_file)
            writer.writerow(CSV_HEADERS)

            with multiprocessing.Pool() as pool:
                results = pool.imap(process_dxf_file, dxf_files)
                for file_number, rows in enumerate(results, start=1):
                    dxf_path = dxf_files[file_number - 1]
                    logging.info(
                        "[%d/%d] Processing %s",
                        file_number,
                        len(dxf_files),
                        dxf_path.relative_to(folder),
                    )
                    rows_from_file = 0
                    for row in rows:
                        writer.writerow(row)
                        all_rows.append(row)
                        row_count += 1
                        rows_from_file += 1
                    if rows_from_file == 0:
                        empty_dxf_files.append(dxf_path)

        logging.info("Wrote %d pipe rows to %s", row_count, output_path)
        if empty_dxf_files:
            print("\nFiles that did not contribute to the CSV:")
            for path in empty_dxf_files:
                print(str(path.relative_to(folder)))
    else:
        logging.info("No DXF files found to process.")

    # Append to consolidated MTO Excel table (or update formulas / summary row)
    append_to_mto_excel(all_rows, mto_excel_path)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    main()
